# -*- coding: utf-8 -*-
"""
ETL: แปลงไฟล์ออเดอร์ของลูกค้า (จาก Google Sheets) → กรอกลงชีต 'order' ของ WAYTHAI-FORYOU_Template.xlsx

วิธีใช้:
  1) ใน Google Sheets:  File → Download → Comma-separated values (.csv)   หรือ  Microsoft Excel (.xlsx)
  2) วางไฟล์ที่ได้ไว้โฟลเดอร์เดียวกับสคริปต์ (หรือระบุ path เป็น argument)
  3) ตรวจ/แก้ COLMAP ด้านล่างให้ตรงกับ "หัวคอลัมน์จริง" ในชีตของคุณ
  4) ติดตั้งไลบรารี (ครั้งเดียว):   pip install pandas openpyxl
  5) รัน:   python etl_to_template.py  orders_export.csv
  6) ได้ไฟล์ WAYTHAI-FORYOU_Template_filled.xlsx → อัปโหลดในแท็บ "นำเข้าข้อมูล"

หมายเหตุ:
  - กรณี "สั่งซัพ" (outsourced) ให้ใส่ยอดจ่ายซัพในคอลัมน์ค่าไม้ (cost_wood) ได้
  - ระบบเว็บจะรวมต้นทุน = ไม้+สี+แรง+ส่ง และคำนวณกำไรผลิตเอง/สั่งซัพ + ค่าสี ให้เอง
"""
import sys, os, datetime
import pandas as pd
from openpyxl import load_workbook

SRC_FILE = sys.argv[1] if len(sys.argv) > 1 else "orders_export.csv"
TEMPLATE = os.path.join(os.path.dirname(__file__), "..", "WAYTHAI-FORYOU_Template.xlsx")
OUT      = os.path.join(os.path.dirname(__file__), "..", "WAYTHAI-FORYOU_Template_filled.xlsx")

# ฟิลด์ในชีต order  →  รายชื่อหัวคอลัมน์ที่อาจเจอในไฟล์ลูกค้า (แก้/เพิ่มได้ตามจริง)
COLMAP = {
    "external_id":    ["รหัสออเดอร์", "เลขที่ออเดอร์", "Order ID", "No", "ลำดับ"],
    "customer_name":  ["ชื่อลูกค้า", "ลูกค้า", "Customer"],
    "phone":          ["เบอร์โทร", "เบอร์", "โทร", "Phone"],
    "description":    ["รายละเอียด", "รายการ", "สินค้า", "ประตู", "งาน"],
    "quantity":       ["จำนวนบาน", "จำนวน", "บาน", "Qty"],
    "production_type":["แหล่งผลิต", "ผลิตเอง/สั่งซัพ", "ผลิต", "Type"],
    "issue_date":     ["วันที่รับงาน", "วันที่", "Date"],
    "due_date":       ["วันนัดส่ง", "กำหนดส่ง", "Due"],
    "total_price":    ["ราคารวม", "ราคา", "ยอดขาย", "Price"],
    "cost_wood":      ["ค่าไม้", "ไม้", "ยอดจ่ายซัพ", "จ่ายซัพ", "Wood"],
    "cost_paint":     ["ค่าสี", "สี", "Paint"],
    "cost_labor":     ["ค่าแรง", "แรง", "Labor"],
    "cost_shipping":  ["ค่าส่ง", "ส่ง", "ขนส่ง", "Shipping"],
    "deposit":        ["มัดจำ", "Deposit"],
    "status":         ["สถานะ", "Status"],
}
# ค่าในคอลัมน์แหล่งผลิต → self / outsourced
PROD_MAP = {"ผลิตเอง": "self", "เอง": "self", "self": "self",
            "สั่งซัพ": "outsourced", "ซัพ": "outsourced", "outsource": "outsourced", "outsourced": "outsourced"}

ORDER_FIELDS = ["external_id", "customer_name", "phone", "description", "quantity", "production_type",
                "issue_date", "due_date", "total_price", "cost_wood", "cost_paint", "cost_labor",
                "cost_shipping", "deposit", "status"]


def find_col(df, names):
    cols = [str(c).strip() for c in df.columns]
    for n in names:                                  # ตรงเป๊ะก่อน
        for i, c in enumerate(cols):
            if c.lower() == str(n).strip().lower():
                return df.columns[i]
    for n in names:                                  # แล้วค่อยจับบางส่วน
        for i, c in enumerate(cols):
            if str(n).strip().lower() in c.lower():
                return df.columns[i]
    return None


def load(path):
    if path.lower().endswith(".csv"):
        return pd.read_csv(path, dtype=str).fillna("")
    return pd.read_excel(path, dtype=str).fillna("")


def norm_date(v):
    v = str(v).strip()
    if not v:
        return ""
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%Y/%m/%d"):
        try:
            return datetime.datetime.strptime(v[:10], fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return v


def num(v):
    s = str(v).replace(",", "").replace("฿", "").strip()
    try:
        return float(s)
    except ValueError:
        return ""


def main():
    if not os.path.exists(SRC_FILE):
        print("ไม่พบไฟล์ต้นทาง:", SRC_FILE)
        print("→ ส่งออก Google Sheets เป็น CSV/XLSX แล้ววางไว้ หรือระบุ path เป็น argument")
        return
    df = load(SRC_FILE)
    cols = {k: find_col(df, v) for k, v in COLMAP.items()}
    matched = {k: str(c) for k, c in cols.items() if c is not None}
    print("จับคู่คอลัมน์ได้:", matched)
    missing = [k for k in ("customer_name", "total_price") if not cols.get(k)]
    if missing:
        print("⚠ หาคอลัมน์สำคัญไม่เจอ:", missing, "— โปรดแก้ COLMAP ให้ตรงหัวคอลัมน์จริง")

    wb = load_workbook(TEMPLATE)
    ws = wb["order"]
    r = 3
    for _, row in df.iterrows():
        def g(k):
            return row[cols[k]] if cols.get(k) is not None else ""
        name = str(g("customer_name")).strip()
        price = str(g("total_price")).strip()
        if not name and not price:
            continue
        ext = str(g("external_id")).strip() or f"ORD-{r-2}"
        pt_raw = str(g("production_type")).strip().lower()
        pt = next((PROD_MAP[k] for k in PROD_MAP if k in pt_raw), "self")
        vals = {
            "external_id": ext, "customer_name": name, "phone": str(g("phone")).strip(),
            "description": str(g("description")).strip(), "quantity": num(g("quantity")),
            "production_type": pt, "issue_date": norm_date(g("issue_date")), "due_date": norm_date(g("due_date")),
            "total_price": num(g("total_price")), "cost_wood": num(g("cost_wood")), "cost_paint": num(g("cost_paint")),
            "cost_labor": num(g("cost_labor")), "cost_shipping": num(g("cost_shipping")),
            "deposit": num(g("deposit")), "status": str(g("status")).strip() or "in_progress",
        }
        for ci, f in enumerate(ORDER_FIELDS, start=1):
            ws.cell(row=r, column=ci, value=vals[f])
        r += 1

    wb.save(OUT)
    print(f"✓ เขียน {r-3} ออเดอร์ → {OUT}")
    print("  เปิดเว็บ → แท็บ 'นำเข้าข้อมูล' → อัปโหลดไฟล์นี้ได้เลย")


if __name__ == "__main__":
    main()
